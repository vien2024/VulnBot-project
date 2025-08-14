from typing import List

import tqdm
from langchain_community.document_loaders.unstructured import UnstructuredFileLoader
from utils.log_common import build_logger


logger = build_logger()


class RapidOCRPPTLoader(UnstructuredFileLoader):

    def _get_elements(self) -> List:
        def ppt2text(filepath):
            from io import BytesIO

            import numpy as np
            from PIL import Image
            from pptx import Presentation
            from rag.parsers.ocr import get_ocr

            try:
                ocr = get_ocr()
            except Exception as e:
                logger.error(
                    f"Failed to initialize OCR engine: {e}. OCR will be disabled."
                )
                ocr = None

            prs = Presentation(filepath)
            resp = ""

            def extract_text(shape):
                nonlocal resp
                if shape.has_text_frame:
                    resp += shape.text.strip() + "\n"
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            for paragraph in cell.text_frame.paragraphs:
                                resp += paragraph.text.strip() + "\n"
                if shape.shape_type == 13 and ocr:  # Picture
                    try:
                        image_bytes = BytesIO(shape.image.blob)
                        pil_image = Image.open(image_bytes)
                        result, _ = ocr(np.array(pil_image))
                        if result:
                            ocr_result = [line[1] for line in result]
                            resp += "\n" + "\n".join(ocr_result)
                    except Exception as e:
                        logger.error(
                            f"Error during OCR on pptx image in {filepath}: {e}"
                        )
                elif shape.shape_type == 6:  # Group
                    for child_shape in shape.shapes:
                        extract_text(child_shape)

            b_unit = tqdm.tqdm(
                total=len(prs.slides), desc="RapidOCRPPTLoader slide index: 1"
            )
            # 遍历所有幻灯片
            for slide_number, slide in enumerate(prs.slides, start=1):
                b_unit.set_description(
                    "RapidOCRPPTLoader slide index: {}".format(slide_number)
                )
                b_unit.refresh()
                try:
                    sorted_shapes = sorted(slide.shapes, key=lambda x: (x.top, x.left))
                    for shape in sorted_shapes:
                        extract_text(shape)
                except Exception as e:
                    logger.error(
                        f"Error processing shape on slide {slide_number} in {filepath}: {e}"
                    )
                b_unit.update(1)
            return resp

        text = ppt2text(self.file_path)
        from unstructured.partition.text import partition_text

        return partition_text(text=text, strategy="fast", **self.unstructured_kwargs)
